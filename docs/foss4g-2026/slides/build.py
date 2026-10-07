# FOSS4G Hokkaido 2026 発表スライド: MIERUNE テンプレートの見た目（背景画像・文字位置）を踏襲して組み立てる（issue #47）
#   uv run --no-project --with python-pptx --with lxml python docs/foss4g-2026/slides/build.py
# テンプレート（docs/foss4g-2026/MIERUNE Template.pptx）と出力の .pptx は git 管理外。
# 地図の画面（shots/）は shots.cjs で撮り直せる（手元で API と npm run dev を起動しておく）
import copy, io, sys
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
SHOTS = HERE / 'shots'
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / 'foss4g2026.pptx'
TEMPLATE = HERE.parent / 'MIERUNE Template.pptx'

# テンプレートの色（背景画像から採った青）と文字色
BLUE = RGBColor(0x30, 0x4C, 0xF0)
INDIGO = RGBColor(0x34, 0x1D, 0xFF)
CYAN = RGBColor(0x06, 0xA8, 0xE0)
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x5B, 0x63, 0x70)
TINT = RGBColor(0xEE, 0xF1, 0xFE)
LINE = RGBColor(0xD5, 0xDB, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xB8, 0xBE, 0xCC)

prs = Presentation(str(TEMPLATE))
BLANK = prs.slide_layouts[6]
# 見本スライドの背景画像（1・2: 表紙、3: 章扉、4: 本文（パンくずあり）、5: 本文）
BG = {}
for i, s in enumerate(prs.slides, 1):
    pic = next(sh for sh in s.shapes if sh.name == 'bg')
    BG[i] = pic.image.blob
N_TEMPLATE = len(prs.slides)


def font_ea(run):
    """テンプレートと同じく、和文も +mn-lt で書く"""
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', '+mn-lt')


def add_slide(bg, notes=''):
    s = prs.slides.add_slide(BLANK)
    pic = s.shapes.add_picture(io.BytesIO(BG[bg]), 0, 0, prs.slide_width, prs.slide_height)
    pic.name = 'bg'
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line=1.2, name=None, bullet=False, space_after=0, autofit=False):
    """paras: 文字列の並び。1 段落の中で色や太さを変えるときは [(文字, {opts}), ...]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp):
        bp.remove(el)
    if autofit:
        etree.SubElement(bp, qn('a:spAutoFit'))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(Emu(Inches(0.32))))
            pPr.set('indent', str(-Emu(Inches(0.32))))
            etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
            etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        runs = para if isinstance(para, list) else [para]
        for t, o in ((r, {}) if isinstance(r, str) else r for r in runs):
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.color.rgb = o.get('color', color)
            font_ea(r)
    return tb


def box(s, x, y, w, h, fill=TINT, line=None, radius=0.08, name=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def image(s, path, x, y, w=None, h=None, border=True, name=None):
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
    if name:
        pic.name = name
    return pic


def content(title, crumb=None, notes=''):
    """本文スライド。crumb ありはパンくず付きの背景（見本 4）、なしは見本 5"""
    if crumb:
        s = add_slide(4, notes)
        text(s, 0.79, 0.23, 4.5, 0.4, [crumb], size=14, bold=True, name='breadcrumb')
        text(s, 0.61, 0.95, 12.1, 0.8, [title], size=36, bold=True, name='title')
    else:
        s = add_slide(5, notes)
        text(s, 0.61, 0.6, 12.1, 0.8, [title], size=36, bold=True, name='title')
    return s


TOP = 2.0  # パンくず付きスライドの本文の上端


def stat(s, x, y, w, value, unit, label, value_color=BLUE, size=40, h=1.55):
    box(s, x, y, w, h)
    text(s, x + 0.25, y + 0.18, w - 0.5, 0.75, [[(value, {'size': size, 'bold': True, 'color': value_color}),
                                                  (unit, {'size': 18, 'bold': True, 'color': value_color})]], line=1.0)
    text(s, x + 0.25, y + 0.98, w - 0.5, 0.5, [label], size=14, color=MUTED, line=1.1)


# ---------------------------------------------------------------- 1. 表紙
s = add_slide(1, '自己紹介は 1 行で。すぐデモに入る')
text(s, 0.61, 1.48, 6, 0.44, ['株式会社MIERUNE'], size=20, color=WHITE, name='name')
text(s, 0.61, 2.3, 10.6, 1.9, ['国土数値情報 道路データで作る、', '全国対応の到達圏・経路探索 Web アプリ'],
     size=36, bold=True, color=WHITE, line=1.15, name='title')
text(s, 0.61, 4.45, 8, 0.9, ['FOSS4G Hokkaido 2026　コアデイ', '2026.11.28　トナリエ北広島'], size=20, color=WHITE,
     line=1.4, name='subtitle')

# ---------------------------------------------------------------- 2. つかみ: 会場から 120 分
s = content('この会場から、車で何分？',
            notes='公開版を映して、会場（トナリエ北広島）をタップ。120 分で札幌・千歳・岩見沢まで色が付く。'
                  '駄目なら予備のローカルデモ → 録画の順に切り替える')
image(s, SHOTS / 'venue_120_z10.png', 0.61, 1.65, h=5.25, name='map')
X = 0.61 + 5.25 * 16 / 9 + 0.35
text(s, X, 1.75, 13.333 - X - 0.5, 0.5, ['トナリエ北広島から 120 分'], size=18, bold=True)
text(s, X, 2.3, 13.333 - X - 0.5, 1.0, [[('11.5', {'size': 48, 'bold': True, 'color': BLUE}),
                                         ('万本', {'size': 20, 'bold': True, 'color': BLUE})]], line=1.0)
text(s, X, 3.2, 13.333 - X - 0.5, 0.5, ['の道路を、到達時間で塗り分け'], size=16, color=MUTED)
text(s, X, 4.1, 13.333 - X - 0.5, 1.0, [[('68', {'size': 48, 'bold': True, 'color': BLUE}),
                                         (' ms', {'size': 20, 'bold': True, 'color': BLUE})]], line=1.0)
text(s, X, 5.0, 13.333 - X - 0.5, 0.5, ['探索して返すまで（手元の PC）'], size=16, color=MUTED)

# ---------------------------------------------------------------- 3. つかみ: 480 分
s = content('北海道全体でも、0.2 秒',
            notes='打ち切りを 480 分に。北海道は本州と道路がつながっていないので、ほぼ全域が色づく')
image(s, SHOTS / 'venue_480_z7.png', 0.61, 1.65, h=5.25, name='map')
text(s, X, 1.75, 13.333 - X - 0.5, 0.5, ['同じ地点から 480 分'], size=18, bold=True)
text(s, X, 2.3, 13.333 - X - 0.5, 1.0, [[('26.8', {'size': 48, 'bold': True, 'color': BLUE}),
                                         ('万本', {'size': 20, 'bold': True, 'color': BLUE})]], line=1.0)
text(s, X, 3.2, 13.333 - X - 0.5, 0.5, ['の道路に到達'], size=16, color=MUTED)
text(s, X, 4.1, 13.333 - X - 0.5, 1.0, [[('85', {'size': 48, 'bold': True, 'color': BLUE}),
                                         (' ms', {'size': 20, 'bold': True, 'color': BLUE})]], line=1.0)
text(s, X, 5.0, 13.333 - X - 0.5, 0.8, ['探索して返すまで。', '描画 111 ms と合わせて 0.2 秒'], size=16, color=MUTED, line=1.3)

# ---------------------------------------------------------------- 4. 背景
s = content('道路ネットワークが、オープンデータになった', '背景',
            notes='経路探索の道路網は有償が当たり前だった。2026 年 4 月に国土数値情報の道路データが CC BY 4.0 で公開')
box(s, 0.61, TOP, 4.6, 4.6, fill=BLUE)
text(s, 0.9, TOP + 0.35, 4.0, 0.5, ['国土数値情報 道路データ'], size=18, bold=True, color=WHITE)
text(s, 0.9, TOP + 1.0, 4.0, 1.2, [[('約 2,400', {'size': 54, 'bold': True, 'color': WHITE}),
                                    ('万本', {'size': 24, 'bold': True, 'color': WHITE})]], line=1.0)
text(s, 0.9, TOP + 2.2, 4.0, 0.5, ['全国の道路中心線'], size=16, color=WHITE)
text(s, 0.9, TOP + 3.05, 4.0, 1.2, ['2026 年 4 月公開', 'CC BY 4.0（商用利用も可）'], size=18, bold=True,
     color=WHITE, line=1.4)
text(s, 5.75, TOP + 0.1, 6.9, 4.4, [
    'これまで、経路探索・到達圏分析に使う道路網は有償のもの（DRM など）が一般的だった',
    '国土交通省が国土数値情報の道路データを公開し、誰でも自由に使えるようになった',
    'ただし配られるのはライン（道路中心線）。経路探索にはノード・リンクの形に組み直す必要がある',
], size=20, bullet=True, line=1.3, space_after=14)

# ---------------------------------------------------------------- 5. これまでの取り組み
s = content('これまでの取り組みと、今日の話', '背景',
            notes='MIERUNE JCT #03（2026-06）で変換と応用を紹介した。そのとき展望に挙げた「全国対応 Web API」の続きが今日')
steps = [
    ('1', '変換する', 'ライン形式の道路データを、ノード・リンク形式の道路ネットワークに組み直す手法を作った', False),
    ('2', '応用する', '到達圏分析・公共交通空白地域の抽出・津波避難シミュレーション\n（MIERUNE JCT #03・2026 年 6 月）', False),
    ('3', '全国で使えるように', '地図のどこをタップしても、全国の道路網で到達圏と経路をすぐ返す Web API とアプリ', True),
]
cw, gap = 3.85, 0.27
for i, (n, head, body, now) in enumerate(steps):
    x = 0.61 + i * (cw + gap)
    box(s, x, TOP + 0.3, cw, 3.9, fill=BLUE if now else TINT)
    fg = WHITE if now else INK
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(TOP + 0.6), Inches(0.6), Inches(0.6))
    circ.fill.solid(); circ.fill.fore_color.rgb = WHITE if now else BLUE; circ.line.fill.background()
    circ.shadow.inherit = False
    tf = circ.text_frame; tf.margin_left = tf.margin_right = 0
    r = tf.paragraphs[0].add_run(); r.text = n; r.font.size = Pt(20); r.font.bold = True
    r.font.color.rgb = BLUE if now else WHITE; tf.paragraphs[0].alignment = PP_ALIGN.CENTER; font_ea(r)
    text(s, x + 0.3, TOP + 1.45, cw - 0.6, 0.6, [head], size=22, bold=True, color=fg)
    text(s, x + 0.3, TOP + 2.15, cw - 0.6, 2.0, body.split('\n'), size=16, color=fg if now else MUTED, line=1.35)
    if now:
        text(s, x + 0.3, TOP + 3.65, cw - 0.6, 0.4, ['← 今日の話'], size=14, bold=True, color=WHITE)

# ---------------------------------------------------------------- 6. 作ったもの: 道路網
s = content('全国 388 万リンクの道路網', '作ったもの',
            notes='幅員 5.5m 以上に絞った。細街路まで入れると 2,400 万本で、車の到達圏には要らない')
cw = (12.11 - 3 * 0.3) / 4
for i, (v, u, l) in enumerate([('388', '万', 'リンク（全体の 16%）'), ('371', '万', 'ノード'),
                               ('37.5', '万 km', '総延長'), ('48', '秒', '全国の道路網の生成')]):
    stat(s, 0.61 + i * (cw + 0.3), TOP + 0.1, cw, v, u, l, size=44, h=1.7)
text(s, 0.61, TOP + 2.2, 12.1, 2.6, [
    [('幅員 5.5m 以上に絞る', {'bold': True}), ('　細街路まで入れると 2,400 万本。車の移動に要る範囲に絞った', {})],
    [('道路 1 本ごとに到達時間を出す', {'bold': True}), ('　メッシュにしないので、どこからでも始められて道の形のまま描ける', {})],
    [('道路種別ごとの固定速度', {'bold': True}), ('　高速 80・国道 35・都道府県道 30・市区町村道 20 km/h', {})],
], size=18, bullet=True, line=1.3, space_after=10)

# ---------------------------------------------------------------- 7. 作ったもの: 構成
s = content('構成', '作ったもの',
            notes='画面・タイル・API を 1 つの Worker の後ろに置いた。DB は持たない')


def node(s, x, y, w, h, head, sub, fill=TINT, fg=INK, subfg=MUTED):
    box(s, x, y, w, h, fill=fill)
    text(s, x + 0.2, y + 0.15, w - 0.4, 0.5, [head], size=18, bold=True, color=fg)
    text(s, x + 0.2, y + 0.65, w - 0.4, h - 0.75, sub, size=14, color=subfg, line=1.3)


def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = BLUE
    c.line.width = Pt(2.25)
    ln = c.line._get_or_add_ln()
    etree.SubElement(ln, qn('a:tailEnd')).set('type', 'triangle')
    return c


node(s, 0.61, TOP + 1.2, 3.0, 2.2, 'ブラウザ', ['SvelteKit', 'MapLibre GL JS', 'PMTiles'])
node(s, 4.4, TOP + 1.2, 2.9, 2.2, 'Cloudflare Worker', ['パスで振り分け', '回数制限・キャッシュ'], fill=BLUE, fg=WHITE, subfg=WHITE)
node(s, 8.1, TOP + 0.0, 4.6, 1.3, '画面（Static Assets）', ['ビルド済みの HTML・JS'])
node(s, 8.1, TOP + 1.6, 4.6, 1.3, '道路タイル（R2）', ['PMTiles 1 ファイル 146 MB・Range で読む'])
node(s, 8.1, TOP + 3.2, 4.6, 1.5, 'API（Containers）', ['FastAPI ＋ SciPy・0.5 vCPU / 4 GiB', '道路網はイメージに焼き込み・DB なし'])
arrow(s, 3.61, TOP + 2.3, 4.4, TOP + 2.3)
arrow(s, 7.3, TOP + 2.0, 8.1, TOP + 0.65)
arrow(s, 7.3, TOP + 2.3, 8.1, TOP + 2.25)
arrow(s, 7.3, TOP + 2.6, 8.1, TOP + 3.95)
text(s, 0.61, TOP + 3.75, 7.2, 0.9, ['同じオリジンの /api・/tiles・それ以外 を 1 つの Worker で振り分ける'], size=14, color=MUTED)

# ---------------------------------------------------------------- 8. 章扉
s = add_slide(3, 'ここから工夫を 3 つ。探索・描画・公開')
text(s, 0.96, 3.9, 11.56, 2.6, ['全国規模で速く動かす', '3 つの工夫'], size=60, bold=True, color=WHITE,
     align=PP_ALIGN.RIGHT, line=1.15, name='title')

# ---------------------------------------------------------------- 9. 工夫① 探索: 比較
s = content('pgRouting と同じ結果のまま、10〜27 倍速い', '工夫① 探索',
            notes='同じ道路網で pgRouting と SciPy（メモリ上のグラフ）を比べた。到達ノード数もコストも完全一致')
cd = CategoryChartData()
cd.categories = ['到達圏 30 分', '到達圏 480 分', '経路 東京→大阪']
cd.add_series('pgRouting（PostgreSQL）', (315, 3400, 3324))
cd.add_series('SciPy（メモリ上のグラフ）', (17, 216, 267))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.61), Inches(TOP), Inches(8.0), Inches(4.7), cd)
gf.name = 'chart'
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(14)
ch.font.size = Pt(14)
ch.font.color.rgb = INK
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = -10
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '#,##0" ms"'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(14)
dl.font.bold = True
for ser, col in zip(plot.series, (GRAY, BLUE)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = 4200
ca = ch.category_axis
ca.reverse_order = True
ca.format.line.fill.background()
ca.tick_labels.font.size = Pt(14)
box(s, 9.0, TOP + 0.3, 3.7, 4.1)
text(s, 9.25, TOP + 0.55, 3.2, 0.5, ['結果は完全一致'], size=18, bold=True)
text(s, 9.25, TOP + 1.1, 3.2, 1.9, ['到達したノードも、コストも同じ。', '違うのは速さだけ'], size=16, color=MUTED, line=1.35)
text(s, 9.25, TOP + 2.7, 3.2, 1.0, [[('×12〜19', {'size': 36, 'bold': True, 'color': BLUE})]], line=1.0)
text(s, 9.25, TOP + 3.4, 3.2, 0.7, ['左の 3 つの比。打ち切り 30〜480 分を通すと 10〜27 倍'], size=14, color=MUTED, line=1.3)

# ---------------------------------------------------------------- 10. 工夫① 探索: 置き場所
s = content('差はアルゴリズムではなく、グラフの置き場所', '工夫① 探索',
            notes='どちらもダイクストラ法。pgRouting は SQL でエッジを毎回読むので、近場を探すだけでも時間がほぼ一定')
for i, (head, lines, hl) in enumerate([
    ('pgRouting', ['呼ぶたびに、SQL でエッジのテーブルを読む', '打ち切っても先に読むので、範囲を変えても時間がほぼ一定', 'データの更新は UPDATE ですぐ反映'], False),
    ('SciPy（採用）', ['起動時に一度だけ道路網をメモリに載せる', '打ち切った分だけ速い（30 分なら 17 ms）', '更新は再起動。今回はデータが年 1 回なので困らない'], True),
]):
    x = 0.61 + i * 6.2
    box(s, x, TOP + 0.1, 5.9, 3.15, fill=BLUE if hl else TINT)
    text(s, x + 0.3, TOP + 0.3, 5.3, 0.5, [head], size=22, bold=True, color=WHITE if hl else INK)
    text(s, x + 0.3, TOP + 0.95, 5.3, 2.2, lines, size=17, color=WHITE if hl else INK, bullet=True, line=1.25,
         space_after=8)
box(s, 0.61, TOP + 3.55, 12.1, 1.15, fill=WHITE, line=LINE)
text(s, 0.9, TOP + 3.68, 11.6, 0.95, [
    [('経路も、毎回全国を探さない', {'bold': True}), ('　直線距離から打ち切りを見積もり、届かなければ倍に広げる', {})],
    [('0.5 vCPU で 東京→横浜 0.98 → ', {'color': MUTED}), ('0.04 秒', {'bold': True, 'color': BLUE})]],
    size=17, line=1.35)

# ---------------------------------------------------------------- 11. 工夫① 探索: 3 エンジン
s = content('既存の経路エンジンを使えばいいのでは？', '工夫① 探索',
            notes='OSRM・GraphHopper・Valhalla に同じ道路網（東京 23 区）を入れ、ランダム 100 組で経路のリンク列が一致するか比べた')
text(s, 0.61, TOP, 12.1, 0.9, ['同じ道路網（東京 23 区・10.6 万リンク）を 3 つのエンジンに入れ、ランダム 100 組で経路を比べた'],
     size=18, color=MUTED)
rows = [('組み合わせ', '同じ経路', '所要の差 平均', '最大', '重なり'),
        ('OSRM – GraphHopper', '63 組', '1 秒', '5 秒', '95%'),
        ('OSRM – Valhalla', '18 組', '25 秒', '140 秒', '71%'),
        ('GraphHopper – Valhalla', '17 組', '24 秒', '140 秒', '71%')]
tbl = s.shapes.add_table(len(rows), 5, Inches(0.61), Inches(TOP + 0.75), Inches(12.1), Inches(2.6))
tbl.name = 'table'
t = tbl.table
t.first_row = True
widths = [4.3, 2.0, 2.2, 1.8, 1.8]
for j, wd in enumerate(widths):
    t.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    t.rows[i].height = Inches(0.65)
    for j, v in enumerate(row):
        c = t.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = BLUE if i == 0 else (TINT if i == 1 else WHITE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = Inches(0.2)
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        c.margin_right = Inches(0.25)
        r = p.add_run(); r.text = v
        r.font.size = Pt(18); r.font.bold = (i == 0 or (i == 1 and j == 1))
        r.font.color.rgb = WHITE if i == 0 else (BLUE if (i == 1 and j == 1) else INK)
        font_ea(r)
text(s, 0.61, TOP + 3.65, 12.1, 1.0, [
    [('3 つとも同じ経路は 15 組。', {'bold': True}), ('外れるのは Valhalla が絡む組み合わせ', {})]], size=18)

# ---------------------------------------------------------------- 12. 工夫① 探索: 厳密さ
s = content('違いを生むのは、言語ではなく「厳密さ」', '工夫① 探索',
            notes='OSRM と GraphHopper は厳密解なので、違いは所要が同じ経路のどれを選ぶかだけ。'
                  'Valhalla は自分の物差しで測っても OSRM の経路の方が速い組が 82 組中 47 組')
for i, (head, sub, lines, hl) in enumerate([
    ('OSRM・GraphHopper', '厳密解（CH・MLD）', ['違いは、所要が同じ経路のどれを選ぶかだけ', '所要の差は最大 5 秒（重みの丸め）'], False),
    ('Valhalla', '速さのために厳密さを手放す', ['階層の打ち切り・双方向 A* の打ち切り', [('自分の物差しでも最短でない', {'bold': True}), ('経路を返す（82 組中 47 組）', {})]], False),
]):
    x = 0.61 + i * 6.2
    box(s, x, TOP + 0.1, 5.9, 2.85)
    text(s, x + 0.3, TOP + 0.3, 5.3, 0.5, [head], size=22, bold=True)
    text(s, x + 0.3, TOP + 0.85, 5.3, 0.45, [sub], size=16, bold=True, color=BLUE)
    text(s, x + 0.3, TOP + 1.4, 5.3, 1.5, lines, size=17, bullet=True, line=1.25, space_after=6)
box(s, 0.61, TOP + 3.25, 12.1, 1.45, fill=BLUE)
text(s, 0.9, TOP + 3.4, 11.6, 1.2, [
    [('到達圏は「全ノードまでの正確な所要」が要る', {'bold': True, 'size': 20})],
    ['→ 厳密なダイクストラ法を SciPy で回す今の形が合っている']], size=17, color=WHITE, line=1.4)

# ---------------------------------------------------------------- 13. 工夫② 描画: 仕組み
s = content('道路はベクタタイル、色は feature-state', '工夫② 描画',
            notes='タイルは道路の形だけを持つ。API はリンク ID と到達時間だけを返し、画面で feature-state に入れて塗る')
image(s, SHOTS / 'venue_120_z12.png', 0.61, TOP, h=4.65, name='map')
X2 = 0.61 + 4.65 * 16 / 9 + 0.4
text(s, X2, TOP + 0.05, 13.333 - X2 - 0.45, 4.6, [
    [('タイルは形だけ', {'bold': True})], ['全国の道路を PMTiles 1 ファイルに（146 MB）'],
    [('API は ID と時間だけ', {'bold': True})], ['到達したリンク ID と到達時間を返す'],
    [('画面で塗る', {'bold': True})], ['MapLibre GL JS の feature-state に入れて、時間ごとに色を変える'],
], size=17, line=1.3, space_after=4)

# ---------------------------------------------------------------- 14. 工夫② 描画: 塗り方
s = content('数十万本を一度に塗ると、画面が止まる', '工夫② 描画',
            notes='feature-state を一度に変えると、タイル枚数 × 本数の再計算が 1 フレームに集まる。'
                  '読み込み済みタイルの道路だけを、太い道から 1 フレーム 4,000 本ずつ塗る')
cd = CategoryChartData()
cd.categories = ['最初の色が出るまで', '最も長く止まったフレーム']
cd.add_series('一度に塗る', (5.8, 4.3))
cd.add_series('太い道から少しずつ', (1.0, 0.24))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.61), Inches(TOP), Inches(6.3), Inches(4.7), cd)
gf.name = 'chart'
ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(14)
ch.font.size = Pt(14); ch.font.color.rgb = INK
plot = ch.plots[0]; plot.gap_width = 70; plot.overlap = -10
plot.has_data_labels = True
dl = plot.data_labels; dl.number_format = '0.0#" 秒"'; dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(14); dl.font.bold = True
for ser, col in zip(plot.series, (GRAY, BLUE)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
va = ch.value_axis; va.visible = False; va.has_major_gridlines = False; va.maximum_scale = 7
ch.category_axis.format.line.color.rgb = LINE
ch.category_axis.tick_labels.font.size = Pt(14)
text(s, 7.3, TOP + 0.05, 5.4, 0.4, ['スマホ相当（CPU 4 倍遅い・9 Mbps）・東京 z11'], size=14, color=MUTED)
text(s, 7.3, TOP + 0.6, 5.4, 4.0, [
    [('原因', {'bold': True})], ['feature-state を変えると、読み込んだタイルごとに塗り直しが走る（タイル枚数 × 本数）'],
    [('対策', {'bold': True})], ['読み込み済みタイルの道路だけを、太い道から 1 フレーム 4,000 本ずつ塗る'],
    ['広域のズームは高速・国道だけのタイルにする'],
], size=17, line=1.3, space_after=6)

# ---------------------------------------------------------------- 15. 工夫③ 公開: 数字
s = content('DB を持たず、Cloudflare の 1 台に載せる', '工夫③ 公開',
            notes='0.5 vCPU・4 GiB のコンテナ 1 台。前処理で読むだけのファイルにし、返す形もバイナリにした')
cw = (12.11 - 0.3) / 2
for i, (head, before, after, u, how) in enumerate([
    ('起動（道路網の読み込み）', '44', '10', '秒', '前処理で、読むだけの配列ファイル（.npy）にしておく'),
    ('常駐メモリ', '2.35', '1.02', 'GiB', '読むだけなので、組み立て用の一時メモリが要らない'),
    ('到達圏の転送量（120 分）', '3.6', '0.95', 'MB', 'JSON をやめ、ID の差分と時間のバイナリで返す'),
    ('道路タイルの読み込み', '0.2', '0.09', '秒', '読んだ範囲を Worker の Cache API に置く（無料）'),
]):
    x = 0.61 + (i % 2) * (cw + 0.3)
    y = TOP + (i // 2) * 2.35
    box(s, x, y, cw, 2.1)
    text(s, x + 0.3, y + 0.18, cw - 0.6, 0.45, [head], size=16, bold=True, color=MUTED)
    text(s, x + 0.3, y + 0.62, cw - 0.6, 0.85, [[(before, {'size': 26, 'bold': True, 'color': GRAY}),
                                                (' → ', {'size': 26, 'bold': True, 'color': GRAY}),
                                                (after, {'size': 40, 'bold': True, 'color': BLUE}),
                                                (' ' + u, {'size': 18, 'bold': True, 'color': BLUE})]], line=1.0)
    text(s, x + 0.3, y + 1.48, cw - 0.6, 0.5, [how], size=15, color=INK)

# ---------------------------------------------------------------- 16. 工夫③ 公開: 費用と当日
s = content('使われていなければ寝る。今日は起こしておく', '工夫③ 公開',
            notes='寝ている間は課金されない。今日は台数を増やし、寝ないようにしてある')
for i, (head, lines) in enumerate([
    ('ふだん', ['コンテナ 1 台（0.5 vCPU・4 GiB）', '10 分使われないと寝る。寝ている間は課金なし', '起きるまで約 20 秒']),
    ('同時アクセス', ['重い計算は同時に 2 本まで（メモリを守る）', '端末ごと・IP ごとに回数制限', '100 本一斉でもエラー 0（手元の計測）']),
    ('今日（FOSS4G）', ['台数を増やして振り分ける', '寝ないようにして、発表前に全台を起こす', '終わったら 1 台に戻す']),
]):
    x = 0.61 + i * (3.85 + 0.27)
    hl = i == 2
    box(s, x, TOP + 0.1, 3.85, 4.5, fill=BLUE if hl else TINT)
    text(s, x + 0.3, TOP + 0.35, 3.25, 0.9, [head], size=20, bold=True, color=WHITE if hl else INK)
    text(s, x + 0.3, TOP + 1.25, 3.25, 3.2, lines, size=16, bullet=True, line=1.3, space_after=10,
         color=WHITE if hl else INK)

# ---------------------------------------------------------------- 17. 割り切りと今後
s = content('割り切りと今後', 'まとめ',
            notes='精度を下げる割り切りであって欠陥ではない、と言う')
text(s, 0.61, TOP, 7.0, 0.5, ['割り切り'], size=20, bold=True)
cw = (7.4 - 0.3) / 2
for i, (head, sub) in enumerate([('車のみ', '徒歩なら細街路こそ経路になる'), ('一方通行は考えない', 'データに属性がない'),
                                 ('道路種別ごとの固定速度', '速度・渋滞の属性はない'), ('コンテナの拠点', '日本に固定できず、応答に約 0.2 秒上乗せ')]):
    x = 0.61 + (i % 2) * (cw + 0.3)
    y = TOP + 0.65 + (i // 2) * 1.95
    box(s, x, y, cw, 1.7)
    text(s, x + 0.25, y + 0.2, cw - 0.5, 0.5, [head], size=18, bold=True)
    text(s, x + 0.25, y + 0.8, cw - 0.5, 0.8, [sub], size=15, color=MUTED)
text(s, 8.4, TOP, 4.3, 0.5, ['今後'], size=20, bold=True)
box(s, 8.4, TOP + 0.65, 4.3, 3.6, fill=BLUE)
text(s, 8.65, TOP + 0.85, 3.8, 3.2, [[('経路の階層化', {'bold': True, 'size': 20})],
                                     ['遠距離の経路を速くする（東京→大阪で今 0.9 秒）'],
                                     [('始点の寄せ方の改善', {'bold': True, 'size': 20})],
                                     ['タップした所から近い道路に、より自然に寄せる']], size=16, color=WHITE, line=1.4,
     space_after=10)

# ---------------------------------------------------------------- 18. まとめ
s = content('まとめ', notes='伝えたいことは 1 つ')
text(s, 0.61, 1.6, 12.1, 1.4, [[('オープンデータの道路網だけで、', {'size': 30, 'bold': True})],
                                [('全国の到達圏をスマホで約 1 秒', {'size': 30, 'bold': True, 'color': BLUE})]], line=1.15)
cw = (12.11 - 2 * 0.3) / 3
for i, (head, body) in enumerate([('どこに置くか', '道路網はメモリに。\npgRouting と同じ結果で 10〜27 倍'),
                                  ('何を送るか', 'タイルは形だけ、API は ID と時間だけ。\n転送量は約 4 分の 1'),
                                  ('何を塗るか', '見えている道路を、太い道から。\nスマホで最初の色まで約 1 秒')]):
    x = 0.61 + i * (cw + 0.3)
    box(s, x, 3.2, cw, 2.9)
    text(s, x + 0.3, 3.45, cw - 0.6, 0.6, [head], size=24, bold=True, color=BLUE)
    text(s, x + 0.3, 4.2, cw - 0.6, 1.8, body.split('\n'), size=17, line=1.35)
text(s, 0.61, 6.35, 12.1, 0.5, ['効いたのは、アルゴリズムよりこの 3 つ'], size=16, color=MUTED)

# ---------------------------------------------------------------- 19. 触ってみてください
s = add_slide(2, 'QR を出したまま質疑へ')
text(s, 0.61, 1.48, 6, 0.44, ['触ってみてください'], size=20, color=WHITE, name='name')
text(s, 0.61, 2.3, 7.4, 1.9, ['会場から始まる', '到達圏・経路探索'], size=48, bold=True, color=WHITE, line=1.15, name='title')
text(s, 0.61, 4.35, 7.4, 0.5, ['ksj-route-search.shi-works-worker.workers.dev'], size=18, color=WHITE)
text(s, 0.61, 5.6, 7.6, 1.0, ['道路データ: 国土数値情報 道路データ（国土交通省）CC BY 4.0', '測量法に基づく国土地理院長承認（使用）R 8JHs 85',
                              '背景地図: 国土地理院 最適化ベクトルタイル'], size=12, color=WHITE, line=1.3)
box(s, 8.75, 1.45, 3.9, 3.9, fill=WHITE, radius=0.06)
s.shapes.add_picture(str(SHOTS / 'qr.png'), Inches(8.95), Inches(1.65), Inches(3.5), Inches(3.5)).name = 'qr'
text(s, 8.75, 5.5, 3.9, 0.5, ['スマホで読み取り'], size=16, color=WHITE, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 見本スライドを外す
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[:N_TEMPLATE]:
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

prs.save(str(OUT))
print('saved', OUT, len(prs.slides), 'slides')
